
"""Import inventory from CSV or Excel (.xlsx).

Preview without saving:
    python -m scripts.import_inventory inventory.csv

Apply changes:
    python -m scripts.import_inventory inventory.csv --apply

Excel:
    python -m scripts.import_inventory inventory.xlsx --apply
"""

import argparse
import csv
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.database.connection import get_session_factory
from app.database.models import Product


REQUIRED_COLUMNS = {
    "name",
    "description",
    "price",
    "stock_quantity",
    "category",
    "active",
}

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_PRICE = Decimal("99999999.99")


class ImportValidationError(Exception):
    """Raised when the inventory file cannot be read or validated."""


def read_rows(path: Path) -> list[dict[str, object]]:
    """Read inventory rows from CSV or XLSX."""
    if not path.is_file():
        raise ImportValidationError(f"File not found: {path}")

    if path.stat().st_size > MAX_FILE_BYTES:
        raise ImportValidationError("File exceeds the 10 MB limit.")

    suffix = path.suffix.lower()

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)

            if not reader.fieldnames:
                raise ImportValidationError("CSV has no header row.")

            headers = [
                header.strip().lower() if header else ""
                for header in reader.fieldnames
            ]

            if any(not header for header in headers):
                raise ImportValidationError(
                    "CSV contains an empty column name."
                )

            if len(headers) != len(set(headers)):
                raise ImportValidationError(
                    "CSV contains duplicate column names."
                )

            missing = REQUIRED_COLUMNS - set(headers)
            if missing:
                raise ImportValidationError(
                    f"Missing required columns: {', '.join(sorted(missing))}"
                )

            rows: list[dict[str, object]] = []

            for row_number, raw in enumerate(reader, start=2):
                if row_number > MAX_ROWS + 1:
                    raise ImportValidationError(
                        f"File exceeds the {MAX_ROWS}-row limit."
                    )

                # Extra values indicate that a CSV row has more fields
                # than the header, usually because of a malformed row.
                if None in raw:
                    raise ImportValidationError(
                        f"CSV row {row_number} has extra fields."
                    )

                normalized = {
                    str(key).strip().lower(): value
                    for key, value in raw.items()
                    if key is not None
                }
                rows.append(normalized)

            return rows

    if suffix == ".xlsx":
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)

            if not first:
                raise ImportValidationError(
                    "Excel worksheet is empty."
                )

            headers = [
                str(value).strip().lower()
                if value is not None
                else ""
                for value in first
            ]

            # Ignore trailing empty Excel columns.
            while headers and not headers[-1]:
                headers.pop()

            if not headers or any(not header for header in headers):
                raise ImportValidationError(
                    "Excel contains an empty column name."
                )

            if len(headers) != len(set(headers)):
                raise ImportValidationError(
                    "Excel contains duplicate column names."
                )

            missing = REQUIRED_COLUMNS - set(headers)
            if missing:
                raise ImportValidationError(
                    f"Missing required columns: {', '.join(sorted(missing))}"
                )

            rows: list[dict[str, object]] = []

            for row_number, values in enumerate(iterator, start=2):
                if row_number > MAX_ROWS + 1:
                    raise ImportValidationError(
                        f"File exceeds the {MAX_ROWS}-row limit."
                    )

                values = list(values)

                if len(values) > len(headers):
                    extra_values = values[len(headers):]
                    if any(value is not None for value in extra_values):
                        raise ImportValidationError(
                            f"Excel row {row_number} has extra fields."
                        )

                values = values[:len(headers)]

                # Skip entirely blank spreadsheet rows.
                if all(value is None for value in values):
                    continue

                values.extend([None] * (len(headers) - len(values)))
                rows.append(dict(zip(headers, values)))

            return rows

        finally:
            workbook.close()

    raise ImportValidationError(
        "File must end in .csv or .xlsx."
    )


def cell_text(value: object) -> str:
    """Convert a cell value to trimmed text."""
    return "" if value is None else str(value).strip()


def parse_active(value: object) -> bool:
    """Parse common spreadsheet boolean values."""
    text = cell_text(value).lower()

    if text in {"true", "1", "yes", "y"}:
        return True

    if text in {"false", "0", "no", "n"}:
        return False

    raise ValueError(
        "active must be true/false, yes/no, or 1/0"
    )


def validate_rows(
    rows: list[dict[str, object]],
) -> tuple[list[dict[str, object]], list[str]]:
    """Validate all products before making database changes."""
    valid: list[dict[str, object]] = []
    errors: list[str] = []
    seen_names: set[str] = set()

    for line, row in enumerate(rows, start=2):
        try:
            name = cell_text(row.get("name"))
            description = cell_text(row.get("description"))
            category = cell_text(row.get("category"))

            if not name:
                raise ValueError("name is required")

            if len(name) > 200:
                raise ValueError(
                    "name must be at most 200 characters"
                )

            if len(description) > 100_000:
                raise ValueError("description is too long")

            if not category:
                raise ValueError("category is required")

            if len(category) > 80:
                raise ValueError(
                    "category must be at most 80 characters"
                )

            key = name.casefold()

            if key in seen_names:
                raise ValueError(
                    f"duplicate product name in file: {name}"
                )

            seen_names.add(key)

            try:
                price = Decimal(cell_text(row.get("price")))
            except (InvalidOperation, ValueError):
                raise ValueError(
                    "price must be a valid decimal number"
                ) from None

            if (
                not price.is_finite()
                or price < 0
                or price > MAX_PRICE
            ):
                raise ValueError(
                    "price must be between 0 and 99999999.99"
                )

            if price.quantize(Decimal("0.01")) != price:
                raise ValueError(
                    "price must have no more than 2 decimal places"
                )

            try:
                stock_decimal = Decimal(
                    cell_text(row.get("stock_quantity"))
                )
            except (InvalidOperation, ValueError):
                raise ValueError(
                    "stock_quantity must be a whole number"
                ) from None

            if (
                not stock_decimal.is_finite()
                or stock_decimal < 0
                or stock_decimal != stock_decimal.to_integral_value()
            ):
                raise ValueError(
                    "stock_quantity must be a non-negative whole number"
                )

            stock = int(stock_decimal)
            active = parse_active(row.get("active"))

            valid.append(
                {
                    "name": name,
                    "description": description,
                    "price": price,
                    "stock_quantity": stock,
                    "category": category,
                    "active": active,
                }
            )

        except (ValueError, TypeError) as exc:
            errors.append(f"Row {line}: {exc}")

    return valid, errors


def import_inventory(path: Path, apply_changes: bool) -> int:
    """Preview or import validated inventory into the configured database."""
    try:
        rows = read_rows(path)
    except (ImportValidationError, OSError, csv.Error) as exc:
        print(f"ERROR: {exc}")
        return 1

    if not rows:
        print("ERROR: No product rows found.")
        return 1

    products, errors = validate_rows(rows)

    if errors:
        print(
            f"Validation failed: {len(errors)} error(s). "
            "Nothing imported."
        )
        for error in errors:
            print(f"  - {error}")
        return 1

    session = get_session_factory()()

    try:
        # Important: begin the transaction before the first query.
        # This avoids SQLAlchemy's already-started transaction error.
        with session.begin():
            names = [item["name"] for item in products]

            existing = session.scalars(
                select(Product).where(Product.name.in_(names))
            ).all()

            by_name = {
                product.name: product
                for product in existing
            }

            created = sum(
                item["name"] not in by_name
                for item in products
            )
            updated = len(products) - created

            print(f"File: {path}")
            print(f"Valid rows: {len(products)}")
            print(f"New products: {created}")
            print(f"Existing products to update: {updated}")

            if not apply_changes:
                print("\nDRY RUN: No changes saved.")
                print("Review the file, then add --apply to import.")

            else:
                for item in products:
                    product = by_name.get(item["name"])

                    if product is None:
                        product = Product(name=item["name"])
                        session.add(product)

                    product.description = item["description"]
                    product.price = item["price"]
                    product.stock_quantity = item["stock_quantity"]
                    product.category = item["category"]
                    product.active = item["active"]

                print("\nImport prepared successfully.")
                print(f"Created: {created}; updated: {updated}")
                print(
                    "WARNING: Existing product stock quantities "
                    "will be replaced by the spreadsheet values."
                )

        # Exiting session.begin() commits successfully, or rolls back
        # automatically if an exception occurs inside the transaction.
        if apply_changes:
            print("\nImport successful.")
        else:
            print("\nPreview completed. No changes were saved.")

        return 0

    except Exception as exc:
        print(
            "ERROR: Import failed; database changes rolled back "
            f"({type(exc).__name__}): {exc}"
        )
        return 1

    finally:
        session.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Import inventory from CSV or XLSX."
    )

    parser.add_argument(
        "file",
        type=Path,
        help="Path to inventory.csv or inventory.xlsx",
    )

    parser.add_argument(
        "--apply",
        action="store_true",
        help="Save changes. Without this flag, preview only.",
    )

    args = parser.parse_args()
    return import_inventory(args.file, args.apply)


if __name__ == "__main__":
    sys.exit(main())